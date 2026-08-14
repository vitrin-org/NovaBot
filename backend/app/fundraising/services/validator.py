# AI Document Validation Service
# Uses config.ai_validation_enabled to switch between mock and real LLM validation

import json
import random
import httpx
from typing import Dict, Optional

from app.config import settings
from app.fundraising.schemas import DocumentTemplate
from app.fundraising.data import TEMPLATES

TEMPLATE_MAP = {t.id: t for t in TEMPLATES}

PROMPTS = {
    "financial-2y": """
شما یک حسابدار رسمی هستید. محتوای فایل‌های مالی را بررسی کنید و فقط JSON زیر را برگردانید:
{"approved": true, "issues": []} یا {"approved": false, "issues": ["مشکل1", "مشکل2"]}

چک‌لیست:
1. آیا دو سال مالی کامل وجود دارد؟
2. آیا امضای حسابرس رسمی موجود است؟
3. آیا ترازنامه و تراز آزمایشی تطابق دارند؟
4. آیا گزارش حسابرس کامل است؟

مثال خروجی:
{"approved": true, "issues": []}
{"approved": false, "issues": ["امضای حسابرس یافت نشد", "سال ۱۴۰۳ ناقص است"]}
""",
    "trial-balance": """
بررسی تراز آزمایشی - فقط JSON برگردانید:
{"approved": true/false, "issues": ["مشکل1", "مشکل2"]}

چک‌لیست:
1. آیا کل و معین ترازند؟
2. آیا تاریخ به روز است؟
3. آیا مهر حسابدار وجود دارد؟
""",
    "credit-report": """
بررسی استعلام اعتباری nics24 - فقط JSON برگردانید:
{"approved": true/false, "issues": ["مشکل1", "مشکل2"]}

چک‌لیست:
1. آیا گواهی رسمی nics24 است؟
2. آیا چک برگشتی رفع سوء‌اثر نشده وجود دارد؟
""",
    "guarantee-letter": """
بررسی پیش‌نویس ضمانت‌نامه - فقط JSON برگردانید:
{"approved": true/false, "issues": ["مشکل1", "مشکل2"]}

چک‌لیست:
1. ذینفع: «شرکت حساب یاری امین ملل» با شناسه 14003497492؟
2. نوع: تعهد پرداخت بدون قید و شرط؟
3. مدت اعتبار: حداقل ۱ سال؟
4. امکان درخواست مجدد قبل از سررسید؟
""",
    "business-plan": """
بررسی مدل کسب و کار - فقط JSON برگردانید:
{"approved": true/false, "issues": ["مشکل1", "مشکل2"]}

چک‌لیست:
1. آیا بوم مدل کسب و کار (BMC) وجود دارد؟
2. تحلیل بازار و رقبا؟
3. پیش‌بینی مالی ۳-۵ ساله؟
""",
    "reg-docs": """
بررسی مدارک ثبتی - فقط JSON برگردانید:
{"approved": true/false, "issues": ["مشکل1", "مشکل2"]}

چک‌لیست:
1. آیا روزنامه رسمی و اساسنامه موجود است؟
2. آیا آگهی آخرین تغییرات شامل است؟
3. آیا شناسه ملی و صاحبان امضا صحیح است؟
""",
    "tech-info": """
بررسی مستندات فنی - فقط JSON برگردانید:
{"approved": true/false, "issues": ["مشکل1", "مشکل2"]}

چک‌لیست:
1. آیا معماری سیستم/نمودار وجود دارد؟
2. آیا tecnologia stack مشخص است؟
3. آیا فاز توسعه و roadmap وجود دارد؟
""",
    "asset-list": """
بررسی لیست دارایی‌ها - فقط JSON برگردانید:
{"approved": true/false, "issues": ["مشکل1", "مشکل2"]}

چک‌لیست:
1. آیا دارایی‌های ثابت/جاری تفکیک شده؟
2. آیا تسهیلات/وام‌ها لیست شده؟
3. آیا اسناد مالکیت موجود است؟
""",
    "hr-info": """
بررسی اطلاعات منابع انسانی - فقط JSON برگردانید:
{"approved": true/false, "issues": ["مشکل1", "مشکل2"]}

چک‌لیست:
1. آیا جدول سهامداری مجموع ۱۰۰٪ می‌شود؟
2. آیا رزومه افراد کلیدی موجود است؟
3. آیا ساختار سازمانی مشخص است؟
""",
    "consent-letter": """
بررسی اجازه‌نامه اعتبارسنجی - فقط JSON برگردانید:
{"approved": true/false, "issues": ["مشکل1", "مشکل2"]}

چک‌لیست:
1. آیا مهر و امضای مجاز موجود است؟
2. آیا با آخرین صاحبان امضا تطابق دارد؟
""",
    "default": """
شما یک بازرس اسناد هستید. محتوای فایل را بررسی کنید و فقط JSON برگردانید:
{"approved": true/false, "issues": []}
"""
}


def extract_text_from_file(file_path: str) -> str:
    """Extract text from PDF, DOCX, XLSX, or text files."""
    text = ""
    try:
        if file_path.lower().endswith(".pdf"):
            import fitz  # PyMuPDF
            doc = fitz.open(file_path)
            text = "\n".join(page.get_text() for page in doc)
        elif file_path.lower().endswith(".docx"):
            from docx import Document as DocxDocument
            doc = DocxDocument(file_path)
            text = "\n".join(p.text for p in doc.paragraphs)
        elif file_path.lower().endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text = "\n".join(
                str(cell.value)
                for ws in wb.worksheets
                for row in ws.iter_rows()
                for cell in row if cell.value
            )
        elif file_path.lower().endswith((".txt", ".md")):
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
        else:
            text = f"[Unsupported file type: {file_path}]"
    except Exception as e:
        text = f"[Error reading file: {e}]"
    return text[:8000]  # Truncate for LLM context


async def call_llm(prompt: str, content: str) -> Dict:
    """Call LLM API and parse JSON response."""
    try:
        response = httpx.post(
            f"{settings.api_base_url}/chat/completions",
            headers={"Authorization": f"Bearer {settings.api_key}"},
            json={
                "model": settings.ai_validation_model or settings.llm_model,
                "messages": [
                    {"role": "system", "content": prompt},
                    {"role": "user", "content": f"متن مدرک:\n{content}"}
                ],
                "temperature": 0.1,
                "max_tokens": 500
            },
            timeout=60
        )
        response.raise_for_status()
        result = response.json()
        ai_response = result["choices"][0]["message"]["content"]

        # Parse JSON from response
        try:
            parsed = json.loads(ai_response)
            return {
                "approved": parsed.get("approved", False),
                "issues": parsed.get("issues", []),
                "raw_response": ai_response
            }
        except json.JSONDecodeError:
            # Try to extract JSON from markdown code blocks
            import re
            json_match = re.search(r'```json\s*(\{.*?\})\s*```', ai_response, re.DOTALL)
            if json_match:
                parsed = json.loads(json_match.group(1))
                return {
                    "approved": parsed.get("approved", False),
                    "issues": parsed.get("issues", []),
                    "raw_response": ai_response
                }
            return {"approved": False, "issues": ["پاسخ AI فرمت JSON نداشت"], "raw_response": ai_response}
    except Exception as e:
        return {"approved": False, "issues": [f"خطا در فراخوانی AI: {e}"], "raw_response": ""}


async def validate_document(template_id: str, file_path: str) -> Dict:
    """
    Validate document using AI (if enabled) or mock (if disabled).
    Returns: {"approved": bool, "issues": List[str], "raw_response": str}
    """
    template = TEMPLATE_MAP.get(template_id)
    if not template:
        return {"approved": False, "issues": ["Template not found"], "raw_response": ""}

    # If AI validation disabled, use mock
    if not settings.ai_validation_enabled:
        import random
        if random.random() > 0.3:
            return {"approved": True, "issues": [], "raw_response": "OK (mock)"}
        else:
            issues = ["مدرک ناقص است", "مهر/امضای معتبر یافت نشد"]
            return {"approved": False, "issues": issues, "raw_response": "Rejected (mock)"}

    # Extract text from file
    content = extract_text_from_file(file_path)
    if not content or content.startswith("[Error") or content.startswith("[Unsupported"):
        return {"approved": False, "issues": ["فایل خالی یا غیرقابل خواندن است"], "raw_response": content}

    # Get prompt for this template
    prompt = PROMPTS.get(template_id, PROMPTS["default"])

    # Call real LLM
    return await call_llm(prompt, content)


def get_template(template_id: str) -> Optional[DocumentTemplate]:
    return TEMPLATE_MAP.get(template_id)


def get_templates_for_platform(platform_id: str):
    from app.fundraising.data import PLATFORMS
    plat = next((p for p in PLATFORMS if p.id == platform_id), None)
    if not plat:
        return []
    return [TEMPLATE_MAP[tid] for tid in plat.required_doc_ids if tid in TEMPLATE_MAP]


def merge_templates(platform_ids: list) -> list:
    """Dedupe templates across multiple platforms."""
    seen = set()
    result = []
    for pid in platform_ids:
        for t in get_templates_for_platform(pid):
            if t.id not in seen:
                seen.add(t.id)
                result.append(t)
    return result